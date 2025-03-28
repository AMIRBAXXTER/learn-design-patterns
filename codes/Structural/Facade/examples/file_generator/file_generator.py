from reportlab.pdfgen import canvas
from xlsxwriter import Workbook
from docx import Document


# Subsystem A: PDF Generator
class PdfGenerator:
    def generate_file(self, filename, content):
        pdf = canvas.Canvas(f"{filename}.pdf")
        pdf.drawString(100, 750, content)
        pdf.save()
        return f"PDF file '{filename}' created successfully."


# Subsystem B: Excel Generator
class ExcelGenerator:
    def generate_file(self, filename, content):
        workbook = Workbook(f"{filename}.xlsx")
        worksheet = workbook.add_worksheet()
        worksheet.write(0, 0, content)
        workbook.close()
        return f"Excel file '{filename}' created successfully."


# Subsystem C: Word Generator
class WordGenerator:
    def generate_file(self, filename, content):
        document = Document()
        document.add_paragraph(content)
        document.save(f"{filename}.docx")
        return f"Word file '{filename}' created successfully."


# Facade: Simple interface to all subsystems
class FileGeneratorFacade:
    def __init__(self):
        self.generators = {
            "pdf": PdfGenerator(),
            "excel": ExcelGenerator(),
            "word": WordGenerator()
        }

    def generate_file(self, file_types: list | tuple, filename: str, content: str) -> list:
        messages = []
        for file_type in [file_type.lower() for file_type in file_types]:
            if file_type in self.generators:
                messages.append(self.generators[file_type].generate_file(filename, content))
            else:
                messages.append(f"File type '{file_type}' is not supported.")
        return messages


# Client Code: Using the Facade
def main():
    facade = FileGeneratorFacade()
    print("\n".join(facade.generate_file(["Pdf"], "test_1", "this is test content")))
    print("\n".join(facade.generate_file(["pdf", "Excel"], "test_2", "this is test_2 content")))
    print("\n".join(facade.generate_file(["pdf", "excel", "Word", "powerpoint"], "test_3", "this is test_3 content")))


if __name__ == "__main__":
    main()
